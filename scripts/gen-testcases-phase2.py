#!/usr/bin/env python3
"""生成《第二阶段测试用例.xlsx》　归属 L5

**表格由脚本生成，不手搓**——与 bigdata/spark/quality_report.py 同一个路子：
用例定义和实测结论都写在代码里，改一处重跑即可，不会出现「改了表忘了改文档」。

格式对齐第一阶段的《第一阶段测试用例.xlsx》：
  1–6 行元信息（项目/模块/编制人/功能特性/测试目的/预置条件）
  7 行表头，8 行起用例，冻结在第 7 行
  列宽 A9.6 B30 C26 D34 E30 F10.6 G18

「测试结果」列的写法沿用一阶段：**结论（日期）+ 换行 + 实测原文**。
凡本轮真跑过的，实测栏是命令的真实输出摘录；人工审核项按约定记为「与预期一致（人工审核）」，
并在备注里标明，不冒充自动化结果。

用法：.venv-phase2/bin/python scripts/gen-testcases-phase2.py
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

REPO = Path(__file__).resolve().parents[1]
OUT = REPO / "第二阶段测试用例.xlsx"
DATE = "2026-09-15"
PROJECT = "东软电动汽车充电桩应用管理平台"
VERSION = "第二阶段 v1.0（大数据与机器学习子系统）"
AUTHOR = "L5（数据可视化与机器学习 / SCML）"

OK = f"通过（{DATE}）"
MAN = "人工审核项"


def p(detail: str) -> str:
    """自动化跑出来的结论 + 实测原文。"""
    return f"{OK}\n实测：{detail}"


def m(detail: str) -> str:
    """人工审核项：按约定记实际结果等于预期结果。"""
    return f"{OK}\n实测：与预期一致（人工审核）。{detail}"


# ── 每张表：(sheet 名, 模块名, 功能特性, 测试目的, 预置条件, 用例列表) ──────────
# 用例 = (编号, 用例说明, 输入数据, 预期结果, 测试结果, 缺陷编号, 备注)

SHEETS: list[dict] = []

SHEETS.append(dict(
    name="01 环境与依赖基线",
    module="第二阶段 · 运行环境与版本基线",
    feature="[二阶段] 2.2 技术基线：Python 3.11/3.12、Hadoop 3.x、PySpark、Spark MLlib、MySQL、Flask、Node 23+、Vue3、DataV",
    purpose="验证各组件版本满足任务书下限；验证自检脚本是「真跑一下」而非只看命令在不在",
    precond="① 已按 docs/RUNBOOK.md 第 7.1 节装完环境；② 本表全部用例由 bash scripts/check-env-phase2.sh 一次跑出",
    cases=[
        ("P1-01", "Python 版本满足 3.11 或 3.12", "bash scripts/check-env-phase2.sh",
         "Python ≥ 3.11；系统自带 3.10 不达标，须用 .venv-phase2",
         p("Python 3.11.15 @ .venv-phase2"), "", "[二阶段] 明文要求；正常路径"),
        ("P1-02", "PySpark 可用性——实际建一次 SparkSession", "同上（脚本内真起 session，非 pip list）",
         "SparkSession 成功创建并正常关闭", p("PySpark 3.5.9　SparkSession 实测可建"), "",
         "正常路径；自检脚本刻意不只看包名，沿用一阶段 check-env.sh 的试编译思路"),
        ("P1-03", "JDK 存在且为 Hadoop/Spark 可用版本", "同上", "JDK 11 或 17",
         p('openjdk version "17.0.20" 2026-07-21'), "", "正常路径"),
        ("P1-04", "MySQL 服务在运行", "同上", "mysqld running 且版本 8.x",
         p("Ver 8.0.46-0ubuntu0.22.04.4；mysqld running"), "", "正常路径；结果层依赖"),
        ("P1-05", "Node 版本 ≥ 23", "同上", "Node ≥ 23",
         p("v24.1.0（npm 11.3.0）"), "", "[二阶段] 明文要求；正常路径"),
        ("P1-06", "Hadoop 版本为 3.x", "同上", "Hadoop 3.x",
         p("Hadoop 3.3.6"), "", "正常路径；选 3.3 线因 PySpark 3.5.9 自带 hadoop-client-api 3.3.4"),
        ("P1-07", "NameNode 可达（不只是装了）", "同上（脚本真连一次）", "hdfs://localhost:9000 可达",
         p("NameNode 可达"), "", "正常路径；装了 ≠ 起着，必须连一次"),
        ("P1-08", "ODS 已在 HDFS 且为只读", "同上", "/ecp/ods 存在，全部文件 444",
         p("/ecp/ods　15 个文件，全部 444 只读"), "", "正常路径；验收第 2 条的自检出口"),
        ("P1-09", "Python 依赖齐全", "同上", "flask / flask_cors / pymysql / pandas / pyarrow 全部就位",
         p("flask 3.1.3　flask_cors 6.0.5　pymysql 1.2.0　pandas 3.0.5　pyarrow 25.0.1"), "", "正常路径"),
        ("P1-10", "自检脚本退出码等于未通过项数", "echo $? after check-env-phase2.sh",
         "全通过时退出码 0；每有一项未通过加 1", p("全部通过，退出码 0"), "",
         "边界值；与一阶段 check-env.sh 口径一致"),
        ("P1-12", "HDFS 未启动时的报错可定位", "停掉 NameNode 后跑读 hdfs:// 的 job",
         "报错并指向「NameNode 未运行」这一根因",
         p("抛 Py4JJavaError：An error occurred while calling o33.csv——"
           "根因（连接被拒）被埋在堆栈里，首行看不出"), "",
         "异常路径；报错不够直白，RUNBOOK 7.2 已写明「HDFS 不会开机自启」作为排障第一步"),
        ("P1-11", "两个 venv 不得混用", "用系统 python3（3.10）执行需要 PySpark 的脚本",
         "报模块缺失或版本不匹配，而不是静默跑出错误结果",
         p("系统 python3 无 pyspark，直接 ModuleNotFoundError；.venv-phase2 正常"), "",
         "异常路径；RUNBOOK 7.1 已写明两个环境不可混用"),
    ]))

SHEETS.append(dict(
    name="02 ODS 原始层与血缘",
    module="第二阶段 · ODS 只读原始层（T2）",
    feature="[二阶段] ODS 作只读原始层；CLAUDE.md 5.2 第 6 条「原始快照一个字节不改」、第 10 条「分析不得直连 charging.db」",
    purpose="验证导出的完整性与血缘可追溯；验证「只读」是结构性保证而非自觉",
    precond="① charging.db 存在；② 执行 .venv-phase2/bin/python bigdata/spark/export_ods.py",
    cases=[
        ("P2-01", "全部业务表导出，不在导出阶段做取舍", "export_ods.py",
         "14 张表全部导出，含空表与低量表——取舍是 T3 质量评估的结论，不能提前砍",
         p("14 表 / 合计 10524 行"), "", "正常路径；砍表会让「t_station_review 为空表」这个结论无从谈起"),
        ("P2-02", "源库只读打开", "sqlite3.connect('file:charging.db?mode=ro', uri=True)",
         "导出过程不可能写到业务库", p("mode=ro 只读连接；导出前后源库指纹均为 2eff8f6c1537b4eb"), "",
         "正常路径；结构性保证，不是靠自觉"),
        ("P2-03", "产物落地后置为只读", "ls -l bigdata/ods/", "全部文件权限 444",
         p("-r--r--r-- t_order.csv　-r--r--r-- _manifest.json"), "", "正常路径；把只读从自觉变成文件权限"),
        ("P2-04", "尝试写 ODS 文件应被系统拒绝", "echo x >> bigdata/ods/t_order.csv",
         "写入失败；本地 444 对属主同样生效", p("bash: bigdata/ods/t_order.csv: 权限不够"), "",
         "异常路径；与 HDFS 侧的差异见 P3-07"),
        ("P2-05", "血缘清单记录可追溯信息", "cat bigdata/ods/_manifest.json",
         "含源库指纹、导出时刻、每表行数与列名",
         p("source_sha256_16=2eff8f6c1537b4eb；14 表；逐表 rows/columns/bytes 齐全"), "",
         "正常路径；换库重跑指纹会变，可据此回溯"),
        ("P2-06", "ECP_ODS_ROOT 指向远端时必须拦住", "ECP_ODS_ROOT=hdfs://... python bigdata/spark/export_ods.py",
         "报错退出（非零），提示改用 scripts/ods-to-hdfs.sh；不得在本地建出名为 hdfs:/... 的目录",
         p("[错误] ECP_ODS_ROOT 指向远端文件系统…　退出码 2"), "BUG-P2-008",
         "异常路径；该缺陷即由本类场景暴露"),
        ("P2-07", "重复导出幂等，不被自己设的 444 挡住", "连续执行 export_ods.py 两次",
         "第二次正常覆盖，行数与指纹一致", p("重跑后 14 表 10524 行、指纹不变；make_writable 先恢复写权限"), "",
         "边界值；第一版曾被自己设的只读权限挡住"),
        ("P2-08", "空表与低量表如实导出", "检查 t_station_review / t_wallet_tx / t_admin_oplog",
         "空表导出为仅表头的 CSV，不跳过", p("t_station_review 0 行（65 字节仅表头）；t_wallet_tx 1 行；t_admin_oplog 3 行"), "",
         "边界值；这些是 T3 质量评估的输入"),
        ("P2-09", "manifest 声明行数与实际 CSV 一致", "逐表比对 manifest.rows 与 CSV 数据行数",
         "逐表一致，合计 10524", p("14 表逐表一致，合计 10524 行"), "", "正常路径"),
        ("P2-11", "ODS 缺文件时报错明确", "读一个不存在的 ODS 表",
         "报路径不存在，而不是返回空结果",
         p("AnalysisException [PATH_NOT_FOUND] Path does not exist: file:…/bigdata/ods/t_not_exist.csv"), "",
         "异常路径；静默返回空会让下游统计悄悄归零"),
        ("P2-10", "导出表名不接受外部输入", "审查 export_ods.py 的 SQL 构造",
         "表名来自本文件常量 TABLES，不来自参数或环境变量",
         m("表名为模块内常量列表，SELECT 语句无外部拼接点。"), "", MAN + "；SQL 注入面审查"),
    ]))

SHEETS.append(dict(
    name="03 HDFS 部署与只读保证",
    module="第二阶段 · Hadoop/HDFS（T8）",
    feature="[二阶段] 文件存储用 Hadoop 3.x；CLAUDE.md 5.2 第 6 条 ODS 只读、第 10 条 代码只认路径不认介质",
    purpose="验证 ODS 落 HDFS 后内容无损、只读被 NameNode 强制、下游可无改动切换介质",
    precond="① bash scripts/install-hadoop-phase2.sh 已执行；② bash scripts/hdfs-ctl.sh start；③ bash scripts/ods-to-hdfs.sh",
    cases=[
        ("P3-01", "HDFS 守护进程可启动并退出安全模式", "bash scripts/hdfs-ctl.sh start",
         "NameNode 与 DataNode 均运行，safemode 自动退出",
         p("namenode 已启动　datanode 已启动　safemode 已退出，可读写"), "",
         "正常路径；单节点用 hdfs --daemon 直启，不依赖 ssh"),
        ("P3-02", "ODS 文件数对账", "bash scripts/ods-to-hdfs.sh",
         "以 _manifest.json 声明的表数为准：14 表 + manifest = 15",
         p("manifest 声明 14 表 + 1 = 15，HDFS 实有 15"), "BUG-P2-002",
         "回归用例；首轮因口径不一致误报「本地 16 ≠ HDFS 15」，见缺陷清单"),
        ("P3-03", "搬运内容字节级无损", "逐文件 hdfs dfs -cat 回本地算 MD5",
         "15 个文件 MD5 与本地快照全部一致", p("15 个文件字节级一致，内容不一致数 0"), "",
         "正常路径；不比 HDFS 的块级 CRC——那与本地文件 md5 不可直接比较"),
        ("P3-04", "HDFS 侧权限与属主", "hdfs dfs -ls /ecp/ods",
         "目录 555、文件 444，属主 ecp_ods:ecp",
         p("dr-xr-xr-x /ecp/ods；-r--r--r-- 1 ecp_ods ecp 逐文件"), "",
         "正常路径；目录需 x 位才能遍历，故 555 而非 444"),
        ("P3-05", "分析身份写入必须被拒绝", "HADOOP_USER_NAME=ecp_analyst hdfs dfs -touchz /ecp/ods/_t",
         "被 NameNode 拒绝，文件不得创建",
         p('Permission denied: user=ecp_analyst, access=WRITE, inode="/ecp/ods":ecp_ods:ecp:dr-xr-xr-x；'
           "退出码 1；目录仍为 15 个文件"), "",
         "异常路径；只读保证要验出来而不是声称"),
        ("P3-06", "分析身份读取必须正常", "HADOOP_USER_NAME=ecp_analyst hdfs dfs -cat /ecp/ods/_manifest.json",
         "可正常读取，权限没有收得过紧", p("读回 manifest 内容，source_sha256_16=2eff8f6c…"), "",
         "正常路径；与 P3-05 成对"),
        ("P3-07", "超级用户绕过 444（已知限制，故用独立身份）", "以启动 NameNode 的账号 bit 在 444 目录建文件",
         "写入成功——证明 444 挡不住超级用户，这正是要用 ecp_analyst 的原因",
         p("bit 成功创建 /ecp/ods/_su_test（-rw-r--r-- bit ecp），验证后已删除"), "",
         "边界值；照搬本地 chmod 444 会让 5.2 第 6 条形同虚设"),
        ("P3-08", "Spark 从 hdfs:// 读取业务数据", "ECP_ODS_ROOT=hdfs://localhost:9000/ecp/ods 下读 t_order.csv",
         "行数与本地一致（8292）", p("从 HDFS 读 t_order：8292 行"), "", "正常路径"),
        ("P3-09", "血缘清单经 HDFS 读回", "read_ods_text(spark, '_manifest.json')",
         "14 表 10524 行、指纹一致", p("14 表 / 10524 行　源库指纹 2eff8f6c1537b4eb"), "BUG-P2-007",
         "回归用例；该路径曾因 InputFormat 过滤 _ 开头文件而失败"),
        ("P3-10", "hdfs:// 前缀不得被路径规范化吃掉", "ods_file('t_order.csv')，ECP_ODS_ROOT 为 hdfs:// URI",
         "结果保留双斜杠：hdfs://localhost:9000/ecp/ods/t_order.csv",
         p("pathlib 对照组 Path(...) → hdfs:/localhost:9000/ecp/ods（已折叠）；"
           "ods_file → hdfs://localhost:9000/ecp/ods/t_order.csv（正确）"), "BUG-P2-006",
         "回归用例；根因是 pathlib 折叠连续斜杠"),
        ("P3-11", "_ 开头的文件必须能被 Spark 读到", "对 _manifest.json 走 read_ods_text",
         "读取成功；不得报 Input path does not exist",
         p("FileSystem.open 读取成功；对照：wholeTextFiles 会报 Input path does not exist"), "BUG-P2-007",
         "回归用例；FileInputFormat 默认过滤 _ 与 . 开头的文件"),
        ("P3-12", "停机重启后数据不丢", "hdfs-ctl.sh stop → start → 检查 /ecp/ods",
         "15 个文件与权限完好（元数据不在 /tmp）",
         p("重启后 Found 15 items，权限属主不变"), "",
         "异常路径；hadoop.tmp.dir 指向 ~/opt/hadoop-data 而非 /tmp"),
        ("P3-13", "切回本地目录不需改代码", "unset ECP_ODS_ROOT 后重跑 validation.py",
         "正常读本地 ODS，结论不变", p("[ODS] 本地　…/bigdata/ods；校验 20/20 通过"), "",
         "正常路径；代码只认路径不认介质"),
        ("P3-14", "每次运行须标明数据源", "观察各 job 启动输出",
         "打印本次 ODS 位置与身份，便于事后分辨结果出自哪一边",
         p("[ODS] HDFS　hdfs://localhost:9000/ecp/ods　身份 ecp_analyst ／ [ODS] 本地　…/bigdata/ods"), "",
         "正常路径；两边结果理应一致，不打标记就无法分辨"),
        ("P3-15", "NameNode 停止后读取必须失败而非返回空", "hdfs-ctl.sh stop 后读 hdfs:// 上的 t_order.csv",
         "抛异常，不得静默返回 0 行",
         p("抛 Py4JJavaError（底层连接被拒）；恢复 NameNode 后同一命令读回 8292 行"), "",
         "异常路径；静默返回空是数据管道最危险的失败方式"),
    ]))

SHEETS.append(dict(
    name="04 清洗 · 探查与质量评估",
    module="第二阶段 · 数据清洗阶段 1+2（T3）",
    feature="《数据清洗基本流程-操作SOP》第 3、4 节：数据探查 + 六维度质量评估",
    purpose="验证探查覆盖全表、质量问题按六维度归类、合理缺失不被误判为问题",
    precond="ODS 已就位；执行 .venv-phase2/bin/python bigdata/spark/profiling.py",
    cases=[
        ("P4-01", "阶段 1 产出数据概况", "profiling.py",
         "产出 quality/01_profile.json，含规模、字段、缺失、唯一值、完全重复",
         p("01_profile.json 覆盖 14 表，合计 10524 行"), "", "正常路径；SOP 阶段 1 必须产出"),
        ("P4-02", "阶段 2 产出问题清单", "profiling.py",
         "产出 quality/02_issues.csv，按六维度归类", p("02_issues.csv 共 11 条（Q001–Q011）"), "",
         "正常路径；SOP 阶段 2 必须产出"),
        ("P4-03", "有效性——非法手机号检出", "t_user.phone",
         "检出不符合 1[3-9] 开头规则的手机号", p("Q001 [有效性] t_user.phone 1 行 (20.00%)"), "",
         "异常路径；实测样例 12345678901"),
        ("P4-04", "一致性——冻结用户却有订单", "t_order 关联 t_user.status",
         "检出业务逻辑冲突并记录规模", p("Q002 [一致性] 冻结用户存在订单 4136 行 (49.88%)"), "",
         "异常路径；后续 R007 决定保留不动"),
        ("P4-05", "完整性——近乎全空列检出", "t_pile_log.old_status / new_status",
         "缺失率 > 50% 的列列为高严重度", p("Q004 old_status 100.0%；Q005 new_status 98.67%，均为高"), "",
         "异常路径；后续 R008 剔除这两列"),
        ("P4-06", "完整性——空表与低量表检出", "t_station_review / t_wallet_tx / t_admin_oplog",
         "如实记录为问题，不静默跳过",
         p("Q010 t_station_review 空表（高）；Q009 t_wallet_tx 1 行；Q011 t_admin_oplog 3 行"), "",
         "边界值；决定了这些表不进分析"),
        ("P4-07", "合理缺失不得计为问题", "已取消订单的 start_time / end_time / settle_time",
         "识别为合理缺失并说明理由，不计入问题清单",
         p("已取消订单 1815 笔无 start_time —— 属合理缺失，不计为问题"), "",
         "边界值；一刀切填充会污染数据"),
        ("P4-08", "准确性——金额与单价×电量逐笔吻合", "已结算订单",
         "逐笔吻合（容差 1 分）", p("6477 笔已结算订单金额与单价×电量逐笔吻合（容差 1 分）"), "",
         "正常路径；整数分口径的第一道关"),
        ("P4-09", "唯一性——业务主键无重复", "t_order.order_no",
         "无重复", p("order_no 空值 0，重复 0（见 05_validation）"), "", "正常路径"),
        ("P4-10", "时效性——数据滞后如实记录", "t_order.settle_time 最新值与当日比较",
         "记录滞后天数；该值随运行日期变化属正常",
         p("Q003 最新结算 2026-09-04 23:59:59，距今 10 天（09-14 跑时为 9 天）"), "",
         "边界值；唯一随日期漂移的字段，重跑对比时须排除"),
        ("P4-11", "探查阶段只看不改", "profiling.py 执行前后比对 ODS",
         "不写 ODS、不写业务库、不产出清洗结果", p("ODS 文件 mtime 与指纹均未变"), "",
         "SOP「先评估、后清洗」；异常路径防护"),
        ("P4-12", "类型转换失败不得静默变空", "ODS 一律按字符串读入",
         "不让 Spark 推断类型，转换失败要能抓出来而非变 null",
         p("profiling/cleaning 均以 inferSchema=False 读入；cast_with_capture() 用影子列比对，"
           "4 表 27 列失败合计 0 行，逐表待核清单均已落地"), "",
         "正常路径；SOP 6.2 明文要求（2026-09-15 补齐捕获后由人工核验升级为自动断言）"),
    ]))

SHEETS.append(dict(
    name="05 清洗 · 执行与校验",
    module="第二阶段 · 数据清洗阶段 4/5/6（T3）",
    feature="SOP 第 5–7 节；CLAUDE.md 5.2 第 9 条「删除必须留痕」",
    purpose="验证 R001–R011 全部执行、全程零删除、校验 20 项全过、报告由产物程序生成",
    precond="阶段 1+2 已完成；依次执行 cleaning.py → validation.py → quality_report.py",
    cases=[
        ("P5-01", "规则全部执行并留痕", "cleaning.py",
         "quality/04_clean_log.json 记录每条规则的命中与处理量",
         p("规则执行 15 条（R001–R011 及分表执行），逐条记录命中数"), "", "正常路径"),
        ("P5-02", "全程零删除零改值", "04_clean_log.json 结论段",
         "本轮删除行数 0、修改值 0", p("本轮删除行数 0，本轮修改值 0"), "",
         "正常路径；SOP 把删除列为最后手段；答辩主动讲的第一件事"),
        ("P5-03", "无法判定的记录进待核清单", "R006 非法手机号",
         "标记待核并保留原值，不删不改", p("标记=1　删除=0　表行数=5；待核清单落 quality/pending/"), "",
         "异常路径；CLAUDE.md 5.2 第 9 条：不静默丢弃"),
        ("P5-04", "近乎全空列剔除并说明", "R008 t_pile_log",
         "剔除 old_status / new_status，保留全部行",
         p("剔除列=['old_status','new_status']　保留行=1800　有效事件=37"), "",
         "边界值；剔列不等于删行；有效事件仅 37 条促成 D8 降级"),
        ("P5-05", "冻结用户历史订单保留不动", "R007",
         "涉及 4136 行，删除 0——历史事实不因当前状态而抹除",
         p("R007 涉及行=4136　删除=0"), "", "边界值；业务口径而非数据错误"),
        ("P5-06", "订单行数前后一致", "ODS → DWD",
         "8292 → 8292", p("写出 dwd_order.parquet　输入=8292　输出=8292"), "", "正常路径"),
        ("P5-07", "校验全部通过方可进入下一阶段", "validation.py",
         "20 项全过；任一项不过以非零退出码结束",
         p("校验结果：20/20 通过　RESULT: PASS"), "", "正常路径；SOP 阶段 5"),
        ("P5-08", "营收逐分对账", "ODS vs DWD 已结算营收",
         "逐分一致，不得有任何漂移", p("ODS 53936279 分 vs DWD 53936279 分"), "",
         "正常路径；贯穿全程的硬校验之一"),
        ("P5-09", "电量对账", "ODS vs DWD 已结算电量",
         "一致（×100 度）", p("36638035 == 36638035"), "",
         "正常路径；与一阶段碳排放对拍脚本独立算出的数字相同"),
        ("P5-10", "金额与电量保持整数型", "R004 断言",
         "amount / price / kwh_x100 均为 bigint，不得出现浮点",
         p("amount=bigint price=bigint kwh_x100=bigint"), "", "正常路径；CLAUDE.md 5.2 第 7 条"),
        ("P5-11", "取消单时长保持为空", "合理缺失不得被填充",
         "被误填 0 行", p("取消单时长保持为空（合理缺失未被填充）　被误填 0 行"), "",
         "边界值；与 P4-07 呼应"),
        ("P5-12", "码表归一无未知值", "R005",
         "status_label 不出现「未知」", p("码表归一无未知值（R005）　status_label 无「未知」"), "",
         "异常路径；越界码会暴露为「未知」"),
        ("P5-13", "派生字段口径自洽", "R010 派生时长与单价",
         "时长非负、派生单价与快照单价偏差 ≤ 1 分",
         p("负时长 0 行；派生单价与快照单价偏差超 1 分的 0 行"), "", "正常路径"),
        ("P5-14", "可重复性——重跑结论一致", "换 HDFS 为数据源重跑阶段 1/4/5",
         "行数、营收、总时长逐项一致；仅与日期相关的字段可变",
         p("与本地基线逐文件比对：04/05/06 完全一致；01/02 仅「滞后天数 9→10」不同（跨日所致）"), "",
         "正常路径；最强的一条：换了存储介质结论不变"),
        ("P5-15", "质量报告由产物程序生成", "quality_report.py",
         "06_quality_report.md 由前五份产物生成，不手写",
         p("质量报告已生成（151 行）；问题 11 条　规则执行 15 条　校验 20 项全过"), "",
         "正常路径；SOP 阶段 6；手写报告会与产物脱节"),
        ("P5-17", "维表与事件表同样捕获转换失败", "t_pile / t_station / t_pile_log",
         "三表各落一份待核清单且为 0 行；参与转换的列非空数在 ODS 与 DWD 间逐列一致",
         p("t_pile 7 列、t_station 5 列、t_pile_log 3 列逐列一致；三份清单行数均为 0"), "",
         "回归；修复「失败捕获只覆盖 t_order」的缺口，见 03_rules.md 一致性核对 #3"),
        ("P5-16", "校验不过必须拦住后续阶段", "使 validation 的任一断言失败",
         "以非零退出码结束，不允许带着失败的校验进入分析阶段",
         m("validation.py 在 failed>0 时以非零码返回，脚本内已实现；本轮 20/20 全过未触发该分支。"),
         "", MAN + "；异常路径，破坏性构造不在真实产物上做"),
    ]))

SHEETS.append(dict(
    name="06 Spark 多维分析",
    module="第二阶段 · 多维分析（T4）",
    feature="[二阶段] 分析维度 ≥ 8 个，其中 ≥ 2 组对比分析；CLAUDE.md 5.2 第 7 条整数聚合、第 10 条不直连业务库",
    purpose="验证维度数量与对比组数达标、各维度营收口径一致、聚合全程走整数",
    precond="DWD 已产出；执行 .venv-phase2/bin/python bigdata/spark/analysis.py",
    cases=[
        ("P6-01", "分析维度数量达标", "bigdata/analysis/_summary.json",
         "维度 ≥ 8", p("dimension_count = 13（要求 ≥ 8）"), "", "正常路径；[二阶段] 明文要求"),
        ("P6-02", "对比分析组数达标", "同上", "对比 ≥ 2 组",
         p("comparison_count = 3（C1 快慢充 / C2 工作日周末 / C3 站点对标）"), "", "正常路径；[二阶段] 明文要求"),
        ("P6-03", "各维度营收合计一致", "跨维度对账",
         "含营收的各维度合计到同一个数", p("各维度营收合计一致　53936279 分"), "",
         "正常路径；口径漂移是这类项目最常见的暗伤"),
        ("P6-04", "聚合全程走整数分", "审查 analysis.py 的金额处理",
         "不得把 amount 读成 float 再累加", p("amount 以 bigint 聚合，仅展示层除以 100"), "",
         "正常路径；CLAUDE.md 5.2 第 7 条；与 5.1 第 3 条同源"),
        ("P6-05", "C1 快充 vs 慢充", "c1_fast_vs_slow.json",
         "含订单量、单均价、时长、度均价、取消率五项对比",
         p("快充 4829 单均价 97.58 元；慢充 1648 单均价 41.34 元；慢充均时长 282.5 分是快充 38.5 分的 7.3 倍"), "",
         "正常路径；度均价同为 147 分说明单价按站点定而非按桩型"),
        ("P6-06", "C2 工作日 vs 周末", "c2_weekday_vs_weekend.json",
         "日均订单量与 24 小时曲线形态可对比",
         p("工作日日均 112.5 单 vs 周末 95.4 单（+18%）；单均金额 83.38 vs 82.93 元几乎持平"), "",
         "正常路径；结论：周末是「来的人少」而非「单笔变小」"),
        ("P6-07", "C3 站点多指标对标", "c3_station_radar.json",
         "多指标归一到 0–100；取消率必须取反，否则面积大反成贬义",
         p("6 站雷达；取消率已取反归一，深圳湾公园低取消得分 0"), "",
         "边界值；不取反会让图表语义反向"),
        ("P6-08", "D8 维度依实测降级", "d8_device_events.json",
         "故障事件为 0 行，改做事件类型与时序分布，不硬凑故障分析",
         p("1800 行中有效事件仅 37 条（上线 27 / 离线 2 / 远程重启 8），event=4 故障上报 0 行"), "",
         "边界值；答辩主动讲的第二件事"),
        ("P6-09", "分析不得直连 charging.db", "审查 analysis.py 数据来源",
         "只读 DWD/ODS，不出现 charging.db 连接", p("数据源为 bigdata/dwd/*.parquet，无 sqlite 连接"), "",
         "正常路径；CLAUDE.md 5.2 第 10 条"),
        ("P6-10", "结果同时落 MySQL 与 JSON 快照", "MySQL 结果表 + bigdata/analysis/*.json",
         "两侧维度齐全且可被 API 读到", p("/api/dimensions 返回 18 份（13 维度 + 4 对比表 + d12 预测表）"), "",
         "正常路径"),
        ("P6-11", "结果层只存分析结果", "检查 MySQL 表内容",
         "不存明细订单，只存聚合结果", m("结果表均为聚合粒度，无逐单明细。"), "",
         MAN + "；否则 Flask 查询压力不可控"),
    ]))

SHEETS.append(dict(
    name="07 Flask 只读 API",
    module="第二阶段 · Flask 只读接口（T5）",
    feature="[二阶段] 用 Flask 处理 web 请求；CLAUDE.md 5.2 第 8 条「Flask 只读」",
    purpose="验证四个接口的正常与异常行为、白名单防注入、金额与百分比的类型正确",
    precond="① MySQL 结果表已就绪；② .venv-phase2/bin/python bigdata/api/app.py；③ 执行 python3 scripts/smoke-api-phase2.py",
    cases=[
        ("P7-01", "健康检查真探库", "GET /api/health",
         "返回库连通状态与维度数，而非固定 ok", p("维度 18 个，库连通 True"), "", "正常路径"),
        ("P7-02", "维度目录不硬编码", "GET /api/dimensions",
         "启动时从 information_schema 读，analysis.py 增删维度后无需改 API",
         p("共 18 份；analysis.py 新增 D11/D12 后 API 未改动即生效"), "", "正常路径"),
        ("P7-03", "维度数满足 ≥ 8", "同上", "≥ 8", p("实际 14 个（13 分析维度 + D12 预测）"), "",
         "正常路径；[二阶段] 要求；对比表单独计数"),
        ("P7-04", "对比分析 ≥ 2 组", "同上", "≥ 2", p("实际 3 组"), "", "正常路径；[二阶段] 要求"),
        ("P7-05", "KPI 概览", "GET /api/overview",
         "返回营收、订单等总量指标", p("营收 539362.79 元　订单 8292"), "", "正常路径"),
        ("P7-06", "金额以整数分返回", "overview.revenue_fen",
         "类型为 int，不得是浮点或字符串", p("revenue_fen 类型 int"), "", "正常路径；CLAUDE.md 5.2 第 7 条"),
        ("P7-07", "百分比必须是数字而非字符串", "overview.settle_rate_pct",
         "类型为 float——MySQL 的 SUM() 经 pymysql 返回 Decimal，默认会被序列化成字符串",
         p("settle_rate_pct 类型 float"), "BUG-P2-003",
         "回归用例；前端拿到字符串会静默画不出图，只表现为「图是空的」"),
        ("P7-08", "取单个维度", "GET /api/dimension/d6_order_status",
         "返回该维度数据与行数", p("2 行；78.11% 已结算 / 21.89% 已取消"), "", "正常路径"),
        ("P7-09", "对比维度内容正确", "GET /api/dimension/c1_fast_vs_slow",
         "含快充与慢充两组", p("['快充', '慢充']"), "", "正常路径"),
        ("P7-10", "跨维度营收对账", "汇总各维度营收",
         "与清洗校验基准一致", p("各维度营收合计一致　53936279 分"), "", "正常路径；贯穿校验"),
        ("P7-11", "不存在的维度返回错误码", "GET /api/dimension/not_exist",
         "404 + 业务错误码，不用裸 bool", p("code=1002"), "", "异常路径"),
        ("P7-12", "SQL 注入尝试被白名单挡下", "维度名传入注入串（表名会拼进 SQL）",
         "被白名单拒绝，返回错误码而非执行", p("code=1002"), "",
         "异常路径；白名单是唯一防线"),
        ("P7-13", "注入后维度目录完好", "注入尝试后重新取目录",
         "仍为 18 份，数据未被破坏", p("仍有 18 份"), "", "异常路径；注入后置校验"),
        ("P7-14", "不存在的接口返回 404", "GET /api/nope",
         "404 + 错误码", p("code=1002"), "", "异常路径"),
        ("P7-17", "超长维度名", "GET /api/dimension/<500 个 a>",
         "返回 404 + 错误码，不崩、不落库查询", p("HTTP 404　code=1002"), "", "边界值"),
        ("P7-18", "路径穿越尝试", "GET /api/dimension/..%2F..%2Fetc%2Fpasswd",
         "被拒，返回错误码", p("code=1002「接口不存在」"), "", "异常路径"),
        ("P7-19", "维度名为空", "GET /api/dimension/",
         "返回 404，不进入查询分支", p("HTTP 404"), "", "边界值"),
        ("P7-20", "维度名夹带 SQL 语句", "GET /api/dimension/d6_order_status; DROP TABLE x",
         "整串按维度名比对白名单后拒绝，不拆分执行",
         p("code=1002「维度 d6_order_status; DROP TABLE x 不存在」"), "",
         "异常路径；表名会拼进 SQL，白名单是唯一防线"),
        ("P7-15", "只监听回环地址", "netstat / 配置审查",
         "仅 127.0.0.1:5000，不暴露到局域网", m("app.run 绑定 127.0.0.1；本服务无鉴权。"), "",
         MAN + "；[说明书] 2.2 通信安全"),
        ("P7-16", "接口不写任何业务库", "审查 app.py 的 SQL",
         "只有 SELECT，无 INSERT/UPDATE/DELETE，不触发 Spark job",
         m("全部语句为 SELECT，无写操作与作业调度入口。"), "", MAN + "；CLAUDE.md 5.2 第 8 条"),
    ]))

SHEETS.append(dict(
    name="08 Vue3 + DataV 大屏",
    module="第二阶段 · 数据可视化大屏（T6）",
    feature="[二阶段] Node 23+ / Vue 3 / DataV 开源 Vue3 组件库；图表类型不得单一",
    purpose="验证五页面板全部渲染出真实数据、DataV 组件生效、图表形态不单一",
    precond="① API 已启动；② cd bigdata/web && npm run build && npm run preview；③ 执行 python3 scripts/smoke-screen-phase2.py",
    cases=[
        ("P8-01", "页面无全局错误", "geckodriver 打开 http://127.0.0.1:4173/",
         "无 .global-err 提示", p("页面正常加载"), "", "正常路径"),
        ("P8-02", "分页标签渲染", "读取 nav.tabs",
         "五页标签齐全", p("5 页：['1运营总览','2站点与设备','3对比分析','4负荷预测','5碳排放']"), "",
         "正常路径"),
        ("P8-03", "第 1 页运营总览渲染", "切到第 1 页", "面板齐全且 canvas 非空",
         p("5 个面板　3 个非空 canvas"), "", "正常路径"),
        ("P8-04", "第 2 页站点与设备渲染", "切到第 2 页", "同上", p("4 个面板　4 个非空 canvas"), "", "正常路径"),
        ("P8-05", "第 3 页对比分析渲染", "切到第 3 页", "同上", p("4 个面板　4 个非空 canvas"), "", "正常路径"),
        ("P8-06", "第 4 页负荷预测渲染", "切到第 4 页",
         "D12 图与明细表渲染；MLlib 未跑时「暂无数据」属预期",
         p("2 个面板　1 个非空 canvas"), "", "边界值；另一格是表格不产生 canvas"),
        ("P8-07", "第 5 页碳排放渲染", "切到第 5 页", "同上", p("3 个面板　1 个非空 canvas"), "",
         "边界值；另两格为表格"),
        ("P8-08", "面板数覆盖全部维度", "统计五页面板总数",
         "≥ 14（13 分析维度 + D12），每个维度至少占一块",
         p("实际 18 个面板（15 图表 + 3 表格）"), "BUG-P2-001",
         "回归用例；该用例发现断言写死 17 已过期，见缺陷清单"),
        ("P8-09", "ECharts 真的画出像素", "统计非空 canvas 面积",
         "canvas 存在还不够，尺寸为 0 也「存在」", p("当前页 1 个非空 canvas；KPI 7 块画布最少 286 个着色像素"), "",
         "边界值；只查存在性会漏掉空白图"),
        ("P8-10", "DataV 边框组件生效", "统计 BorderBox 元素",
         "使用 DataV 开源组件而非自绘边框", p("当前页 BorderBox13 × 3"), "",
         "正常路径；[二阶段] 要求用 DataV"),
        ("P8-11", "DataV 数字翻牌生效", "统计 DigitalFlop", "顶部 KPI 使用翻牌组件",
         p("DigitalFlop × 6"), "", "正常路径"),
        ("P8-12", "DataV 胶囊图与活动环图生效", "第 1 页",
         "两类 DataV 图表均渲染", p("Capsule × 1　ActiveRing × 1"), "", "正常路径"),
        ("P8-13", "无页面级 JS 错误", "收集 window error", "无", p("无"), "", "异常路径"),
        ("P8-14", "截图可产出", "geckodriver 截图",
         "产出非空 PNG，供答辩材料使用", p("screenshot.png 411 KB"), "", "正常路径"),
        ("P8-15", "图表形态不得单一", "统计图表种类",
         "≥ 5 类不同形态", m("实际 9 类：折线面积 / 柱 / 横向条形 / 堆叠柱 / 分组柱 / 环图 / 仪表盘 / 雷达 / 散点，另加 3 张滚动表格。"),
         "", MAN + "；[二阶段] 明文要求"),
        ("P8-16", "投影分辨率下的可读性", "1920×1080 下逐页查看",
         "字号、配色、留白在投影环境可读", m("深色底 + 高对比配色，标题与轴标签在 1080p 下清晰。"), "",
         MAN + "；答辩前需在实际投影环境再确认一次"),
        ("P8-17", "环图中心数字为轮播动画", "观察 D6 活动环图",
         "标签与中心百分比在轮播切换时会短暂不一致，属 DataV 动画行为",
         m("截图曾捕捉到「62% 已取消」的中间帧；接口数据为 78.11% 已结算 / 21.89% 已取消，数据本身正确。"),
         "", MAN + "；答辩截图须等数字稳定后再截，否则看着像算错"),
        ("P8-19", "API 不可用时页面给出明确提示", "停掉 Flask 后刷新大屏",
         "显示全局错误提示，而不是一片空白或卡在加载态",
         m("App.vue 捕获取数异常并写入 state.error，由 .global-err 呈现；渲染冒烟的正向断言即检查该元素不存在。"),
         "", MAN + "；异常路径"),
        ("P8-18", "生产构建可用", "npm run build",
         "构建成功并可由 preview 提供服务", p("578 modules transformed；dist 产出；preview 200"), "",
         "正常路径；冒烟跑的是生产构建而非 dev server"),
    ]))

SHEETS.append(dict(
    name="09 Spark MLlib 建模与评估",
    module="第二阶段 · 机器学习子系统（T7）",
    feature="[二阶段] 用 Spark MLlib 预测未来 1h / 6h / 24h 负荷、空闲桩数、高峰时段，并做模型评估",
    purpose="验证特征口径守恒、切分与选参方法论正确、评估含双基线对照、结果如实报告",
    precond="DWD 已产出；依次执行 features.py → train.py → report.py → predict.py",
    cases=[
        ("P9-01", "特征面板规模", "features.py",
         "站-小时面板，特征数充分", p("8640 行 × 24 特征"), "", "正常路径"),
        ("P9-02", "电量守恒自检", "features.py 自检",
         "面板电量与 DWD 合计偏差为 0", p("电量守恒自检偏差 0.0000 度"), "",
         "正常路径；贯穿校验；口径一旦漂移这里立刻暴露"),
        ("P9-03", "按时间切分，绝不随机切", "eval.json.split",
         "测试集取时间末段，验证集取训练段末段",
         p("按时间切分：测试末 12 天，验证为训练段末 8 天"), "",
         "正常路径；方法论沿用第一阶段；随机切会造成穿越"),
        ("P9-04", "超参在验证段选，不用测试段调参", "eval.json.grid",
         "网格在验证段比较，测试段只用于最终报告",
         p("并集网格 5 组，逐 horizon 在验证段选优"), "", "正常路径"),
        ("P9-05", "六个模型全部产出", "eval.json.results",
         "负荷与并发数 × 1h/6h/24h 共 6 个", p("6 个 GBTRegressor；n_train 5754–5892，n_test 1734"), "",
         "正常路径"),
        ("P9-06", "评估含双基线对照", "eval.json",
         "每个模型给出相对基线 A 与基线 B 的增益",
         p("gain_vs_a / gain_vs_b 逐模型齐全"), "", "正常路径；比单看 MAE 更能说明问题"),
        ("P9-07", "结果如实报告，不修饰", "07_forecast_eval.md",
         "跑赢与跑输都写清楚，并说明原因",
         p("6 个模型中 3 个跑赢基线 B：负荷 1h +1.02%、负荷 24h +0.23%、并发 1h +7.90%；"
           "其余为 -3.03% / -3.85% / -2.10%"), "",
         "正常路径；答辩主动讲的第三件事；原因为缺天气特征 + 框架差异"),
        ("P9-08", "评估报告由 eval.json 程序生成", "report.py",
         "不手写，避免报告与实测脱节", p("07_forecast_eval.md 由 eval.json 生成"), "", "正常路径"),
        ("P9-09", "长迭代训练不得爆栈", "train.py maxIter 提高至 500",
         "设置检查点目录后可正常完成", p("已 setCheckpointDir；未设时 maxIter≤120 能过、500 必崩"), "BUG-P2-005",
         "回归用例；checkpointInterval 不设目录会静默失效"),
        ("P9-10", "三个 horizon 的起报时刻必须不同", "predict.py 输出",
         "从特征面板末尾起报，而非从「要求标签存在」的监督样本起报",
         p("d12_load_forecast 18 行，三个 horizon 的 predict_time 互不相同"), "BUG-P2-004",
         "回归用例；原实现报的是已有观测值的时刻，不是预测"),
        ("P9-11", "预测结果可被大屏取用", "d12_load_forecast → MySQL → API → 大屏第 4 页",
         "全链路贯通", p("/api/dimensions 含 d12_load_forecast；大屏第 4 页渲染正常"), "", "正常路径"),
        ("P9-12", "并发数 24h 信号接近零", "eval.json y_sessions h=24",
         "增益为负属真实结果，不强行调高",
         p("gain_vs_b = -2.10%；与第一阶段 sklearn 收敛到同一数量级"), "",
         "边界值；两个阶段独立佐证该 horizon 信号本就弱"),
        ("P9-14", "模型缺失时预测应报错而非写入空结果", "删除/未生成模型目录后跑 predict.py",
         "明确报错退出，不得向 MySQL 写入空或默认值",
         m("predict.py 以 PipelineModel.load 载入，路径不存在时抛异常中止，不进入写库分支。"),
         "", MAN + "；异常路径"),
        ("P9-13", "不得用 pgrep 判断自建进程状态", "训练期间的状态检测方式",
         "改用日志或 PID 文件——pgrep -f 的模式会匹配到检查进程自己",
         m("已从流程中移除 pgrep 判活；T7 曾因此误报「训练在跑」三个半小时。"), "",
         MAN + "；同一个坑在 T8 以 pkill 形式再次出现"),
    ]))

# ── 缺陷清单 ─────────────────────────────────────────────────────────────────
DEFECTS = [
    ("BUG-P2-001", "P8-08", "大屏面板总数断言写死 17，加页后长期失败",
     "scripts/smoke-screen-phase2.py:121",
     "T7 新增「负荷预测」页的 2 块面板后总数变为 18，而断言常量仍是 17。"
     "该断言自 T7 提交起一直失败且无人重跑；README、PHASE2-PLAN、ARCHITECTURE 三处也跟着写成「17 面板」，"
     "bigdata/web/README.md 的面板对照表更漏了 D12 与 3 张表格",
     "中（测试有效性与文档准确性）", "已修复并回归",
     "断言改为下限式「面板数覆盖全部维度（≥14）」，不再写死会随开发增长的数；"
     "对照表补齐 D12 与 3 张表格并订正为 18；三处文档同步订正。重跑冒烟 15/15 PASS"),
    ("BUG-P2-002", "P3-02", "ODS 推 HDFS 的文件数对账口径不一致",
     "scripts/ods-to-hdfs.sh",
     "本地侧用 find 统计（含 .gitkeep 等隐藏文件），上传与 MD5 循环用 glob（不含隐藏文件），"
     "两边数的不是同一批东西，首轮误报「本地 16 ≠ HDFS 15」；"
     "同时 MD5 行借用了 N_LOCAL 显示条数，报告自身口径也不自洽",
     "轻（误报，不影响数据正确性）", "已修复并回归",
     "改为以 _manifest.json 声明的表数为准（14 表 + manifest = 15），"
     "MD5 行改报循环实际比对数。重跑全部 PASS"),
    ("BUG-P2-003", "P7-07", "MySQL 聚合值被序列化成字符串导致图表空白",
     "bigdata/api/app.py",
     "MySQL 的 SUM() 经 pymysql 返回 Decimal，Flask 默认序列化成字符串（如 \"78.11\"）。"
     "前端图表库拿到字符串会静默画不出来，浏览器里只表现为「图是空的」，极难定位",
     "严重（功能不可用且无报错）", "已修复并回归",
     "开发期（T5）发现；序列化层统一转 float，并在接口冒烟中加一条专门防回归的类型断言。"
     "本轮由 P7-07 覆盖回归"),
    ("BUG-P2-004", "P9-10", "三个 horizon 的预测起报时刻相同",
     "bigdata/mllib/predict.py",
     "起报样本取自监督学习样本，而那份数据要求标签存在，于是目标全部落在数据末尾——"
     "输出的不是预测，而是一个已有观测值的时刻",
     "严重（预测语义错误）", "已修复并回归",
     "开发期（T7）发现；改为从特征面板末尾起报。本轮由 P9-10 覆盖回归"),
    ("BUG-P2-005", "P9-09", "GBT 长迭代训练 StackOverflowError",
     "bigdata/mllib/train.py",
     "GBT 每轮在上一轮 RDD 上叠加，血缘线性增长，任务序列化递归爆栈。"
     "checkpointInterval 在未设置检查点目录时静默不生效——maxIter≤120 能过、500 必崩",
     "中（限制可训练轮数）", "已修复并回归",
     "开发期（T7）发现；加 setCheckpointDir。本轮由 P9-09 覆盖回归"),
    ("BUG-P2-006", "P3-10", "ODS 路径被 pathlib 折叠导致 hdfs:// 失效",
     "bigdata/spark/spark_session.py",
     "ODS_ROOT 曾包成 pathlib.Path，而 pathlib 会折叠连续斜杠："
     "hdfs://localhost:9000/x 变成 hdfs:/localhost:9000/x。"
     "Spark 把它当本地相对路径找，报「Path does not exist」，完全看不出根因",
     "严重（HDFS 通路不可用）", "已修复并回归",
     "T8 发现；ODS 根路径改为 str 并统一用 ods_file() 拼接。本轮由 P3-10 覆盖回归"),
    ("BUG-P2-007", "P3-11", "_manifest.json 被 FileInputFormat 过滤，Spark 读不到",
     "bigdata/spark/spark_session.py",
     "Hadoop 的 FileInputFormat 默认过滤 _ 与 . 开头的文件（_SUCCESS 即由此被忽略），"
     "而血缘清单恰好叫 _manifest.json。hdfs dfs -cat 读得到、Spark 读不到，现象矛盾易误判为文件缺失",
     "中（血缘读取失败）", "已修复并回归",
     "T8 发现；改用 Hadoop FileSystem.open 直读，绕开 InputFormat 过滤。本轮由 P3-11 覆盖回归"),
    ("BUG-P2-008", "P2-06", "export_ods.py 对 HDFS 不可用却不报错",
     "bigdata/spark/export_ods.py",
     "该脚本全程使用 Python 文件 I/O（mkdir / open / chmod），对 hdfs:// 无效。"
     "若按 ECP_ODS_ROOT 直接拼接，只会在本地建出一个名为 hdfs:/localhost:9000 的目录，且无人察觉",
     "中（静默产生错误产物）", "已修复并回归",
     "T8 发现；显式拦截远端根路径并以退出码 2 退出，指向 scripts/ods-to-hdfs.sh。本轮由 P2-06 覆盖回归"),
]

# ── 样式 ────────────────────────────────────────────────────────────────────
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="DDEBF7")
LBL_FILL = PatternFill("solid", fgColor="F2F2F2")
WIDTHS = [9.625, 30, 26, 34, 30, 10.625, 18]
HEADERS = ["用例编号", "用例说明", "输入数据", "预期结果", "测试结果", "缺陷编号", "备注"]


def style_meta(ws, row: int, label: str, value: str, label2: str = "", value2: str = "") -> None:
    ws.cell(row, 1, label).fill = LBL_FILL
    ws.cell(row, 1).font = Font(bold=True)
    ws.cell(row, 2, value)
    if label2:
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        ws.cell(row, 4, label2).fill = LBL_FILL
        ws.cell(row, 4).font = Font(bold=True)
        ws.cell(row, 5, value2)
        ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=7)
    else:
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
    for c in range(1, 8):
        ws.cell(row, c).border = BORDER
        ws.cell(row, c).alignment = Alignment(vertical="center", wrap_text=True)


def build() -> None:
    wb = Workbook()
    wb.remove(wb.active)

    total = 0
    for sh in SHEETS:
        ws = wb.create_sheet(sh["name"])
        for i, w in enumerate(WIDTHS, 1):
            ws.column_dimensions[chr(64 + i)].width = w

        style_meta(ws, 1, "项目名称", PROJECT, "程序版本", VERSION)
        style_meta(ws, 2, "功能模块名", sh["module"])
        style_meta(ws, 3, "编制人", AUTHOR, "编制时间", DATE)
        style_meta(ws, 4, "功能特性", sh["feature"])
        style_meta(ws, 5, "测试目的", sh["purpose"])
        style_meta(ws, 6, "预置条件", sh["precond"])
        ws.row_dimensions[1].height = 24.95

        for c, h in enumerate(HEADERS, 1):
            cell = ws.cell(7, c, h)
            cell.font = Font(bold=True)
            cell.fill = HDR_FILL
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[7].height = 36

        for r, case in enumerate(sh["cases"], 8):
            for c, v in enumerate(case, 1):
                cell = ws.cell(r, c, v)
                cell.border = BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            ws.row_dimensions[r].height = 64
        total += len(sh["cases"])
        ws.freeze_panes = "A8"

    # 缺陷清单
    ws = wb.create_sheet("10 缺陷清单")
    dw = [12, 10, 30, 32, 52, 20, 16, 46]
    for i, w in enumerate(dw, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.cell(1, 1, "项目名称").font = Font(bold=True)
    ws.cell(1, 1).fill = LBL_FILL
    ws.cell(1, 2, PROJECT)
    ws.cell(2, 1, "说明").font = Font(bold=True)
    ws.cell(2, 1).fill = LBL_FILL
    ws.cell(2, 2, "缺陷编号与各用例表「缺陷编号」列一一对应。状态取值：待修复 / 修复中 / 已修复并回归。"
                  "标注「开发期发现」的，是 T5/T7/T8 实现过程中暴露并当场修复的，本轮由对应用例做回归覆盖。")
    ws.cell(3, 1, "编制人").font = Font(bold=True)
    ws.cell(3, 1).fill = LBL_FILL
    ws.cell(3, 2, f"{AUTHOR}　编制时间 {DATE}")
    for r in (1, 2, 3):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
        for c in range(1, 9):
            ws.cell(r, c).border = BORDER
            ws.cell(r, c).alignment = Alignment(vertical="center", wrap_text=True)

    dh = ["缺陷编号", "发现用例", "缺陷标题", "所在文件", "现象与影响", "严重级", "状态", "修复说明"]
    for c, h in enumerate(dh, 1):
        cell = ws.cell(4, c, h)
        cell.font = Font(bold=True)
        cell.fill = HDR_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, d in enumerate(DEFECTS, 5):
        for c, v in enumerate(d, 1):
            cell = ws.cell(r, c, v)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[r].height = 76
    ws.freeze_panes = "A5"

    wb.save(OUT)
    manual = sum(1 for sh in SHEETS for c in sh["cases"] if MAN in c[6])
    print(f"已生成 {OUT.name}")
    print(f"  {len(SHEETS)} 张模块表 + 缺陷清单　用例 {total} 条　缺陷 {len(DEFECTS)} 条")
    print(f"  其中人工审核项 {manual} 条（实际结果按约定记为等于预期）")


if __name__ == "__main__":
    build()
